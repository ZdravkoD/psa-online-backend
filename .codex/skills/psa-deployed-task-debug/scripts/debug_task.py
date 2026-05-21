#!/usr/bin/env python3
"""Fetch deployed-task artifacts and print the next debugging steps."""

from __future__ import annotations

import argparse
import csv
import json
import re
import subprocess
import sys
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path


DEFAULT_API_BASES = [
    "https://psa-online-functions-staging.azurewebsites.net/api",
    "https://psa-online-functions.azurewebsites.net/api",
]
DEFAULT_AZURE_LOGIN_SCOPE = "https://management.core.windows.net//.default"
TASK_ID_RE = re.compile(r"(?P<task_id>[0-9a-fA-F]{24})")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Resolve a deployed PSA task and download its artifacts."
    )
    parser.add_argument(
        "target",
        help="Task-progress URL or raw 24-char task id",
    )
    parser.add_argument(
        "--api-base",
        action="append",
        default=[],
        help="Extra Functions API base(s) to try, for example https://host/api",
    )
    parser.add_argument(
        "--output-dir",
        default="/private/tmp/psa-deployed-task-debug",
        help="Directory where the task folder should be created",
    )
    parser.add_argument(
        "--preview-rows",
        type=int,
        default=12,
        help="Number of input rows to preview",
    )
    parser.add_argument(
        "--skip-downloads",
        action="store_true",
        help="Fetch task metadata only",
    )
    parser.add_argument(
        "--skip-azure-auth-check",
        action="store_true",
        help="Do not validate Azure CLI auth or auto-launch az login",
    )
    parser.add_argument(
        "--azure-login-scope",
        default=DEFAULT_AZURE_LOGIN_SCOPE,
        help="Scope to pass to az login when the Azure session is expired",
    )
    return parser.parse_args()


def extract_task_id(target: str) -> tuple[str, str | None]:
    match = TASK_ID_RE.search(target)
    if not match:
        raise ValueError(f"Could not find a 24-char task id in: {target}")
    task_id = match.group("task_id")
    source_url = target if target.startswith(("http://", "https://")) else None
    return task_id, source_url


def dedupe_keep_order(items: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for item in items:
        normalized = item.rstrip("/")
        if normalized in seen:
            continue
        seen.add(normalized)
        result.append(normalized)
    return result


def build_candidate_api_bases(source_url: str | None, extra_bases: list[str]) -> list[str]:
    candidates: list[str] = []
    if source_url:
        parsed = urllib.parse.urlparse(source_url)
        if parsed.scheme and parsed.netloc:
            candidates.append(f"{parsed.scheme}://{parsed.netloc}/api")
    candidates.extend(extra_bases)
    candidates.extend(DEFAULT_API_BASES)
    return dedupe_keep_order(candidates)


def http_get_bytes(url: str) -> bytes:
    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": "psa-deployed-task-debug/1.0",
            "Accept": "*/*",
        },
    )
    with urllib.request.urlopen(request, timeout=30) as response:
        return response.read()


def resolve_task(task_id: str, api_bases: list[str]) -> tuple[str, dict]:
    errors: list[str] = []
    for api_base in api_bases:
        task_url = f"{api_base}/task/{task_id}"
        try:
            data = json.loads(http_get_bytes(task_url).decode("utf-8"))
            return api_base, data
        except urllib.error.HTTPError as exc:
            errors.append(f"{task_url} -> HTTP {exc.code}")
        except urllib.error.URLError as exc:
            errors.append(f"{task_url} -> URL error: {exc.reason}")
        except json.JSONDecodeError as exc:
            errors.append(f"{task_url} -> invalid JSON: {exc}")
    joined = "\n".join(errors)
    raise RuntimeError(f"Could not resolve task {task_id}.\n{joined}")


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def safe_write_json(path: Path, payload: dict) -> None:
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def download_file(url: str, destination: Path) -> int:
    data = http_get_bytes(url)
    destination.write_bytes(data)
    return len(data)


def maybe_download_input_file(task: dict, api_base: str, task_dir: Path) -> tuple[Path | None, list[str]]:
    notes: list[str] = []
    file_type = task.get("file_type")
    file_data = task.get("file_data")
    file_name = task.get("file_name") or "input"

    if file_type == "json_content" and isinstance(file_data, dict):
        destination = task_dir / "input" / "input.json"
        ensure_dir(destination.parent)
        safe_write_json(destination, file_data)
        notes.append("Saved JSON input content.")
        return destination, notes

    if file_type != "blob_storage_url" or not isinstance(file_data, str):
        notes.append("Task did not contain a downloadable input file.")
        return None, notes

    parsed = urllib.parse.urlparse(file_data)
    blob_name = Path(parsed.path).name
    quoted_blob_name = urllib.parse.quote(blob_name)
    destination = task_dir / "input" / (file_name or blob_name)
    ensure_dir(destination.parent)

    candidate_urls = [file_data]
    if blob_name:
        candidate_urls.append(f"{api_base}/input-file/{quoted_blob_name}")

    errors: list[str] = []
    for candidate_url in dedupe_keep_order(candidate_urls):
        try:
            size = download_file(candidate_url, destination)
            notes.append(f"Downloaded input file from {candidate_url} ({size} bytes).")
            return destination, notes
        except urllib.error.HTTPError as exc:
            errors.append(f"{candidate_url} -> HTTP {exc.code}")
        except urllib.error.URLError as exc:
            errors.append(f"{candidate_url} -> URL error: {exc.reason}")

    notes.append("Input file download failed:")
    notes.extend(errors)
    return None, notes


def download_screenshots(task: dict, task_dir: Path) -> list[str]:
    image_urls = task.get("image_urls") or []
    if not image_urls:
        return ["No screenshot URLs on the task."]

    screenshots_dir = task_dir / "screenshots"
    ensure_dir(screenshots_dir)
    notes: list[str] = []
    for index, image_url in enumerate(image_urls, start=1):
        basename = Path(urllib.parse.urlparse(image_url).path).name or f"screenshot_{index}.png"
        destination = screenshots_dir / f"{index:02d}_{basename}"
        try:
            size = download_file(image_url, destination)
            notes.append(f"Downloaded screenshot {destination.name} ({size} bytes).")
        except urllib.error.HTTPError as exc:
            notes.append(f"Failed to download {image_url}: HTTP {exc.code}")
        except urllib.error.URLError as exc:
            notes.append(f"Failed to download {image_url}: {exc.reason}")
    return notes


def preview_input_file(path: Path, preview_rows: int) -> list[str]:
    lines = [f"Input preview from {path.name}:"]
    suffix = path.suffix.lower()

    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            for idx, row in enumerate(reader, start=1):
                lines.append(f"{idx}: {row}")
                if idx >= preview_rows:
                    break
        return lines

    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError:
            lines.append("openpyxl is not installed in the current interpreter.")
            return lines

        workbook = load_workbook(path, read_only=True, data_only=True)
        lines.append(f"Sheets: {workbook.sheetnames}")
        worksheet = workbook[workbook.sheetnames[0]]
        for idx, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            lines.append(f"{idx}: {row}")
            if idx >= preview_rows:
                break
        return lines

    if suffix == ".xls":
        lines.append("Legacy .xls preview is not implemented. Open it with a spreadsheet tool if needed.")
        return lines

    lines.append("Preview not implemented for this file type.")
    return lines


def _looks_like_azure_auth_failure(output: str) -> bool:
    normalized = output.lower()
    return any(
        marker in normalized
        for marker in [
            "aadsts700082",
            "interactive authentication is needed",
            "please run:\naz login",
            "az login --scope",
            "refresh token has expired",
        ]
    )


def ensure_azure_login(scope: str) -> list[str]:
    notes: list[str] = []
    try:
        status = subprocess.run(
            ["az", "account", "show"],
            capture_output=True,
            text=True,
            check=False,
        )
    except FileNotFoundError:
        return ["Azure CLI is not installed; skipping Azure auth check."]

    if status.returncode == 0:
        notes.append("Azure CLI session is already valid.")
        return notes

    combined_output = "\n".join(
        part.strip() for part in [status.stdout, status.stderr] if part and part.strip()
    )
    if not _looks_like_azure_auth_failure(combined_output):
        notes.append(f"Azure CLI auth check failed with exit code {status.returncode}.")
        if combined_output:
            notes.append(combined_output)
        return notes

    login_command = ["az", "login", "--scope", scope]
    notes.append("Azure CLI session is expired; triggering az login.")

    if not (sys.stdin.isatty() and sys.stdout.isatty()):
        notes.append(
            "This run is non-interactive, so az login cannot complete here. "
            f"Rerun in a TTY: {' '.join(login_command)}"
        )
        return notes

    login_result = subprocess.run(login_command, check=False)
    if login_result.returncode != 0:
        notes.append(f"az login exited with code {login_result.returncode}.")
        return notes

    verify_result = subprocess.run(
        ["az", "account", "show"],
        capture_output=True,
        text=True,
        check=False,
    )
    if verify_result.returncode == 0:
        notes.append("Azure CLI login refreshed successfully.")
    else:
        notes.append("az login ran, but Azure CLI auth still could not be verified.")
    return notes


def detect_environment(api_base: str) -> str:
    lowered = api_base.lower()
    if "staging" in lowered:
        return "staging"
    if "azurewebsites.net" in lowered:
        return "production"
    return "unknown"


def repo_hints(task: dict) -> list[str]:
    haystack = " ".join(
        [
            str(task.get("status", {}).get("message", "")),
            str(task.get("status", {}).get("detailed_error_message", "")),
            " ".join(task.get("distributors") or []),
        ]
    ).lower()

    hints: list[str] = []
    if "opening phoenix login page" in haystack:
        hints.append(
            "Phoenix login bootstrap: scraper/scraper/pharmacy_distributors/phoenix/phoenix.py"
        )
        hints.append(
            "Shared browser navigation/debug context: scraper/scraper/pharmacy_distributors/common/browser_common.py"
        )
    if "prepare_for_order" in haystack or "selecting phoenix client" in haystack:
        hints.append(
            "Phoenix order preparation: scraper/scraper/pharmacy_distributors/phoenix/phoenix.py"
        )
    if "sting" in haystack and "login" in haystack:
        hints.append(
            "Sting login flow: scraper/scraper/pharmacy_distributors/sting/sting.py"
        )
    if "timeout" in haystack:
        hints.append(
            "Check browser startup/navigation retries and the target site's loading state."
        )
    return hints


def print_azure_commands(environment: str) -> None:
    print("\nAzure commands:")
    print("  az account show")
    print("  # If Azure says interactive auth is needed, the helper script will try to trigger:")
    print(f"  az login --scope {DEFAULT_AZURE_LOGIN_SCOPE}")
    print("  az container show --name scraper-container --resource-group psa -o json")
    print("  az container logs --name scraper-container --resource-group psa")
    if environment == "staging":
        print("  az webapp log tail --name psa-online-functions --resource-group psa --slot staging")
    elif environment == "production":
        print("  az webapp log tail --name psa-online-functions --resource-group psa")
    else:
        print("  az webapp log tail --name psa-online-functions --resource-group psa --slot staging")
        print("  az webapp log tail --name psa-online-functions --resource-group psa")
    print("  # Current CI workflows target the same scraper-container resource name;")
    print("  # correlate container logs with the task timestamp before concluding.")


def main() -> int:
    args = parse_args()
    task_id, source_url = extract_task_id(args.target)
    api_bases = build_candidate_api_bases(source_url, args.api_base)
    api_base, task = resolve_task(task_id, api_bases)

    root_dir = Path(args.output_dir)
    task_dir = root_dir / task_id
    ensure_dir(task_dir)
    safe_write_json(task_dir / "task.json", task)

    print(f"Task id: {task_id}")
    print(f"Resolved API base: {api_base}")
    print(f"Task status: {task.get('status', {}).get('status')}")
    print(f"Task message: {task.get('status', {}).get('message')}")
    print(f"Distributors: {task.get('distributors')}")
    print(f"Pharmacy id: {task.get('pharmacy_id')}")
    print(f"Created: {task.get('date_created')}")
    print(f"Updated: {task.get('date_updated')}")
    print(f"Artifacts dir: {task_dir}")

    if not args.skip_azure_auth_check:
        azure_auth_notes = ensure_azure_login(args.azure_login_scope)
        if azure_auth_notes:
            print("\nAzure auth:")
            for note in azure_auth_notes:
                print(f"  - {note}")

    detailed_error = task.get("status", {}).get("detailed_error_message")
    if detailed_error:
        print("\nDetailed error:")
        print(detailed_error)

    hints = repo_hints(task)
    if hints:
        print("\nLikely repo paths:")
        for hint in hints:
            print(f"  - {hint}")

    if args.skip_downloads:
        print_azure_commands(detect_environment(api_base))
        return 0

    screenshot_notes = download_screenshots(task, task_dir)
    print("\nScreenshot downloads:")
    for note in screenshot_notes:
        print(f"  - {note}")

    input_path, input_notes = maybe_download_input_file(task, api_base, task_dir)
    print("\nInput file:")
    for note in input_notes:
        print(f"  - {note}")

    if input_path and input_path.exists():
        print()
        for line in preview_input_file(input_path, args.preview_rows):
            print(line)

    print_azure_commands(detect_environment(api_base))
    return 0


if __name__ == "__main__":
    sys.exit(main())
